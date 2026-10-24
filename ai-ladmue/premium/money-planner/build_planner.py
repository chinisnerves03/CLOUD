from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as CL

FN='Tahoma'
def f(**k): return Font(name=FN,**k)
BLUE=f(size=11,color='0000FF'); BLK=f(size=11); BOLD=f(size=11,bold=True)
H1=f(size=20,bold=True,color='FFFFFF'); H2=f(size=13,bold=True,color='1F3A5F'); GRN=f(size=11,color='008000')
YEL=PatternFill('solid',fgColor='FFF2A8'); HEAD=PatternFill('solid',fgColor='1F6F54'); BAND=PatternFill('solid',fgColor='E8F5EE'); GREY=PatternFill('solid',fgColor='F2F2F2')
thin=Side(style='thin',color='BBBBBB'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
THB='#,##0;(#,##0);-'; THB2='#,##0.00;(#,##0.00);-'; PCT='0.0%'
wb=Workbook()

def title(ws,text,sub=None,width=8):
    ws.sheet_view.showGridLines=False
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=width)
    c=ws.cell(1,1,text);c.font=H1;c.fill=HEAD;c.alignment=Alignment(vertical='center',indent=1)
    ws.row_dimensions[1].height=34
    for i in range(2,width+1): ws.cell(1,i).fill=HEAD
    if sub:
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=width)
        c=ws.cell(2,1,sub);c.font=f(size=10,italic=True,color='555555');c.alignment=Alignment(wrap_text=True,vertical='top',indent=1);ws.row_dimensions[2].height=32
def head(ws,r,vals,c0=1):
    for i,v in enumerate(vals):
        c=ws.cell(r,c0+i,v);c.font=f(size=11,bold=True,color='FFFFFF');c.fill=HEAD;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);c.border=BOX
def inp(c,fmt=None):
    c.font=BLUE;c.fill=YEL;c.border=BOX
    if fmt:c.number_format=fmt
def calc(c,fmt=None,bold=False):
    c.font=BOLD if bold else BLK;c.border=BOX
    if fmt:c.number_format=fmt

# ---------------- 1 เริ่มที่นี่
ws=wb.active;ws.title='เริ่มที่นี่'
title(ws,'แผนเงินของฉัน: เก็บเงิน + ปลดหนี้','ไฟล์นี้คำนวณให้อัตโนมัติ คุณกรอกเฉพาะช่องสีเหลือง ตัวอักษรสีน้ำเงิน',6)
ws.column_dimensions['A'].width=4;ws.column_dimensions['B'].width=90
rows=[('วิธีใช้ 4 ขั้น',H2),
('1) แท็บ "งบรายเดือน": กรอกรายรับและรายจ่ายต่อเดือนของคุณ (ตัวเลขในไฟล์เป็นตัวอย่าง ให้ลบแล้วใส่ของจริง)',BLK),
('2) แท็บ "ปลดหนี้": กรอกหนี้แต่ละก้อน (ยอดคงเหลือ ดอกเบี้ย/ปี ค่างวดขั้นต่ำ) แล้วเลือกวิธีปลดหนี้',BLK),
('3) แท็บ "เป้าหมายเก็บเงิน": ใส่เป้าหมายและจำนวนเดือน ไฟล์จะบอกว่าต้องเก็บเดือนละ/วันละเท่าไหร่',BLK),
('4) แท็บ "คำสั่ง AI": คัดลอกคำสั่งไปวางใน ChatGPT/Claude ให้ช่วยตรวจและปรับแผน (อย่าใส่เลขบัญชีหรือรหัสผ่าน)',BLK),
('',BLK),('สีในไฟล์',H2),
('ช่องสีเหลือง + ตัวอักษรสีน้ำเงิน = ให้คุณกรอก   |   ตัวอักษรสีดำ = สูตรคำนวณ อย่าพิมพ์ทับ   |   สีเขียว = ดึงค่าจากแท็บอื่น',BLK),
('',BLK),('ข้อควรรู้',H2),
('• ไฟล์นี้เป็นเครื่องคำนวณเพื่อวางแผนเบื้องต้น ไม่ใช่คำแนะนำการลงทุน กฎหมาย หรือภาษี',BLK),
('• ผลคำนวณขึ้นกับตัวเลขที่คุณกรอก ดอกเบี้ยจริงและค่าธรรมเนียมของเจ้าหนี้อาจต่างจากที่ไฟล์ประมาณ ให้ตรวจกับใบแจ้งหนี้จริง',BLK),
('• ตารางปลดหนี้จำลองสูงสุด 120 เดือน และรองรับหนี้ 6 ก้อน',BLK),
('• ถ้ามีปัญหาหนี้หนักหรือถูกทวงหนี้ผิดวิธี ควรปรึกษาหน่วยงานที่ให้คำปรึกษาด้านหนี้โดยตรง',BLK),
('• ไม่ได้รวมเงินเฟ้อ ภาษี และความเสี่ยงของผลตอบแทนใดๆ',BLK),
('',BLK),('AI ลัดมือ · @ailadmue',f(size=10,color='777777'))]
for i,(t,fo) in enumerate(rows,start=4):
    c=ws.cell(i,2,t);c.font=fo;c.alignment=Alignment(wrap_text=True,vertical='top')
    if len(t)>95: ws.row_dimensions[i].height=32

# ---------------- 2 งบรายเดือน
bg=wb.create_sheet('งบรายเดือน')
title(bg,'งบรายเดือน','ตัวเลขเป็นตัวอย่าง ให้แทนด้วยของจริงของคุณ (หน่วย: บาท/เดือน)',5)
for col,w in zip('ABCDE',[34,16,16,40,4]):bg.column_dimensions[col].width=w
head(bg,4,['รายรับ','บาท/เดือน','',''])
inc=[('เงินเดือน/รายได้หลัก',45000),('รายได้เสริม',3000),('รายได้อื่นๆ',0)]
for i,(n,v) in enumerate(inc):
    r=5+i;bg.cell(r,1,n).font=BLK;bg.cell(r,1).border=BOX;inp(bg.cell(r,2,v),THB)
bg.cell(8,1,'รวมรายรับ').font=BOLD;calc(bg.cell(8,2,'=SUM(B5:B7)'),THB,True);bg.cell(8,1).border=BOX
head(bg,10,['รายจ่าย','บาท/เดือน','ประเภท','หมายเหตุ'])
exp=[('ที่พัก/ค่าเช่า/ผ่อนบ้าน',8000,'จำเป็น'),('อาหาร',6000,'จำเป็น'),('เดินทาง/น้ำมัน',2500,'จำเป็น'),('ค่าน้ำ ไฟ เน็ต โทรศัพท์',2200,'จำเป็น'),('ประกัน/สุขภาพ',1000,'จำเป็น'),
     ('กิน/เที่ยว/บันเทิง',2500,'อยากได้'),('ช้อปปิ้ง',1500,'อยากได้'),('สมาชิกรายเดือน (Netflix ฯลฯ)',400,'อยากได้'),('อื่นๆ',1000,'อยากได้')]
for i,(n,v,t) in enumerate(exp):
    r=11+i;bg.cell(r,1,n).font=BLK;bg.cell(r,1).border=BOX;inp(bg.cell(r,2,v),THB);inp(bg.cell(r,3,t));bg.cell(r,4).border=BOX
dv=DataValidation(type='list',formula1='"จำเป็น,อยากได้"',allow_blank=False);bg.add_data_validation(dv);dv.add('C11:C19')
r=20;bg.cell(r,1,'ค่างวดขั้นต่ำของหนี้ทั้งหมด (ดึงจากแท็บ ปลดหนี้)').font=BLK;bg.cell(r,1).border=BOX
c=bg.cell(r,2,"=SUM('ปลดหนี้'!D6:D11)");c.font=GRN;c.border=BOX;c.number_format=THB;inp(bg.cell(r,3,'จำเป็น'));bg.cell(r,3).font=BLK;bg.cell(r,3).fill=GREY
bg.cell(21,1,'รวมรายจ่าย').font=BOLD;bg.cell(21,1).border=BOX;calc(bg.cell(21,2,'=SUM(B11:B20)'),THB,True)
bg.cell(23,1,'เงินเหลือต่อเดือน (รายรับ − รายจ่าย)').font=BOLD;bg.cell(23,1).border=BOX;calc(bg.cell(23,2,'=B8-B21'),THB,True)
bg.cell(23,4,'ถ้าติดลบ แปลว่ารายจ่ายมากกว่ารายรับ ต้องปรับก่อน').font=f(size=10,italic=True,color='C0392B')
head(bg,25,['แบ่งเงินที่เหลือ','บาท/เดือน','',''])
bg.cell(26,1,'ใช้ปลดหนี้เพิ่ม (นอกเหนือจากค่างวดขั้นต่ำ)').font=BLK;bg.cell(26,1).border=BOX;inp(bg.cell(26,2,1500),THB)
bg.cell(27,1,'ใช้เก็บเงินตามเป้าหมาย').font=BLK;bg.cell(27,1).border=BOX;inp(bg.cell(27,2,8000),THB)
bg.cell(28,1,'เหลือหลังแบ่งแล้ว').font=BOLD;bg.cell(28,1).border=BOX;calc(bg.cell(28,2,'=B23-B26-B27'),THB,True)
bg.cell(28,4,'=IF(B28<0,"แบ่งเกินกว่าที่มี ลดช่องสีเหลืองด้านบน","โอเค")').font=BOLD
bg.conditional_formatting.add('B23',CellIsRule(operator='lessThan',formula=['0'],font=f(bold=True,color='C0392B')))
bg.conditional_formatting.add('B28',CellIsRule(operator='lessThan',formula=['0'],font=f(bold=True,color='C0392B')))
head(bg,30,['ดูสัดส่วนรายจ่าย','บาท','% ของรายรับ',''])
bg.cell(31,1,'จำเป็น (รวมค่างวดหนี้)').border=BOX;calc(bg.cell(31,2,'=SUMIF(C11:C20,"จำเป็น",B11:B20)'),THB);calc(bg.cell(31,3,'=IF($B$8=0,0,B31/$B$8)'),PCT)
bg.cell(32,1,'อยากได้').border=BOX;calc(bg.cell(32,2,'=SUMIF(C11:C20,"อยากได้",B11:B20)'),THB);calc(bg.cell(32,3,'=IF($B$8=0,0,B32/$B$8)'),PCT)
bg.cell(33,1,'เงินออม/ปลดหนี้เพิ่ม').border=BOX;calc(bg.cell(33,2,'=B26+B27'),THB);calc(bg.cell(33,3,'=IF($B$8=0,0,B33/$B$8)'),PCT)
for r in (31,32,33):bg.cell(r,1).font=BLK
bg.cell(34,1,'หมายเหตุ: สัดส่วนอ้างอิงกฎ 50/30/20 (จำเป็น 50% อยากได้ 30% เก็บ/ปลดหนี้ 20%) เป็นแค่แนวทางทั่วไป ไม่ใช่เกณฑ์ตายตัว').font=f(size=10,italic=True,color='555555')

# ---------------- 3 ปลดหนี้
dp=wb.create_sheet('ปลดหนี้')
title(dp,'แผนปลดหนี้ (จำลองสูงสุด 120 เดือน)','กรอกหนี้สูงสุด 6 ก้อนในช่องสีเหลือง ไฟล์จะเรียงลำดับตามวิธีที่เลือกและจำลองการจ่ายให้ ตัวอย่างในไฟล์เป็นตัวเลขสมมติ',12)
dp.cell(3,1,'วิธีปลดหนี้').font=BOLD;dp.merge_cells('B3:C3');inp(dp['B3'],None);dp['B3']='ดอกเบี้ยสูงสุดก่อน'
dv2=DataValidation(type='list',formula1='"ดอกเบี้ยสูงสุดก่อน,ยอดน้อยสุดก่อน"');dp.add_data_validation(dv2);dv2.add('B3')
dp.cell(3,4,'ดอกเบี้ยสูงสุดก่อน = จ่ายดอกเบี้ยรวมน้อยกว่า | ยอดน้อยสุดก่อน = ปิดก้อนเล็กได้เร็ว ให้กำลังใจ').font=f(size=10,italic=True,color='555555')
head(dp,5,['#','ชื่อหนี้','ยอดคงเหลือ (บาท)','ค่างวดขั้นต่ำ/เดือน','ดอกเบี้ย/ปี','ลำดับที่จ่ายก่อน'])
ex=[('บัตรเครดิต A',45000,2250,0.16),('ผ่อนมือถือ',12000,1500,0.0),('สินเชื่อส่วนบุคคล',60000,2800,0.25),('',0,0,0),('',0,0,0),('',0,0,0)]
for i,(n,b,m,a) in enumerate(ex):
    r=6+i;dp.cell(r,1,i+1).font=BLK;dp.cell(r,1).border=BOX
    inp(dp.cell(r,2,n or None));inp(dp.cell(r,3,b),THB);inp(dp.cell(r,4,m),THB);inp(dp.cell(r,5,a),PCT)
    # rank: avalanche -> higher APR first; snowball -> lower balance first; ignore empty (balance 0) -> last. tie-break by row
    key=f'IF($C{r}<=0,1E9+{r},IF($B$3="ดอกเบี้ยสูงสุดก่อน",1000-$E{r}*1000+{r}/1000,$C{r}+{r}/1000))'
    keys='{}'
    dp.cell(r,6,f'=IF($C{r}<=0,"-",COUNTIF($M$6:$M$11,"<"&M{r})+1)').font=BLK;dp.cell(r,6).border=BOX
    dp.cell(r,13,'='+key).font=f(size=9,color='999999')
dp.cell(5,13,'(คีย์จัดลำดับ)').font=f(size=9,color='999999')
dp.cell(12,2,'รวม').font=BOLD;calc(dp.cell(12,3,'=SUM(C6:C11)'),THB,True);calc(dp.cell(12,4,'=SUM(D6:D11)'),THB,True)
dp.cell(13,2,'เงินเพิ่มปลดหนี้/เดือน (จากแท็บ งบรายเดือน)').font=BLK;c=dp.cell(13,3,"='งบรายเดือน'!B26");c.font=GRN;c.border=BOX;c.number_format=THB
dp.cell(14,2,'งบจ่ายหนี้รวมต่อเดือน (คงที่ทุกเดือน)').font=BOLD;calc(dp.cell(14,3,'=D12+C13'),THB,True)
# helper: sorted debts (priority order) rows 18-23
dp.cell(16,1,'ตารางเรียงตามลำดับที่จ่ายก่อน (ใช้คำนวณ)').font=H2
head(dp,17,['ลำดับ','ชื่อ','ยอดคงเหลือ','ขั้นต่ำ','ดอกเบี้ย/ปี'])
for k in range(6):
    r=18+k;dp.cell(r,1,k+1).font=BLK
    m=f'MATCH({k+1},$F$6:$F$11,0)'
    dp.cell(r,2,f'=IFERROR(INDEX($B$6:$B$11,{m}),"")').font=BLK
    dp.cell(r,3,f'=IFERROR(INDEX($C$6:$C$11,{m}),0)').number_format=THB
    dp.cell(r,4,f'=IFERROR(INDEX($D$6:$D$11,{m}),0)').number_format=THB
    dp.cell(r,5,f'=IFERROR(INDEX($E$6:$E$11,{m}),0)').number_format=PCT
    for cc in range(1,6):dp.cell(r,cc).border=BOX;dp.cell(r,cc).font=BLK if cc==1 or cc==2 else BLK
# summary outputs
dp.cell(25,1,'สรุปผล').font=H2
S0=40; N=120  # simulation start header row 40, data rows 41..160
last=S0+N
dp.cell(26,2,'ปลดหนี้หมดในอีก (เดือน)').font=BLK;calc(dp.cell(26,3,f'=IF(COUNTIF($C$12,0)=1,0,IF(INDEX($B${S0+1}:$B${last},{N})>0.5,"เกิน 120 เดือน",COUNTIF($B${S0+1}:$B${last},">0.5")+1))'),'0',True)
dp.cell(27,2,'ดอกเบี้ยรวมที่จ่ายตลอดแผน (บาท)').font=BLK;calc(dp.cell(27,3,f'=SUM($D${S0+1}:$D${last})'),THB,True)
dp.cell(28,2,'จ่ายรวมทั้งหมด (บาท)').font=BLK;calc(dp.cell(28,3,f'=SUM($E${S0+1}:$E${last})'),THB,True)
dp.cell(29,2,'ตรวจสอบ: ถ้างบจ่ายรวมน้อยกว่าดอกเบี้ยรายเดือน หนี้จะไม่ลด').font=f(size=10,italic=True,color='C0392B')
dp.cell(29,3,f'=IF($C$14<SUMPRODUCT($C$18:$C$23,$E$18:$E$23)/12,"หนี้ไม่ลด ต้องเพิ่มงบ","ปกติ")').font=BOLD
dp.cell(30,2,'ปีและเดือนตัวอย่างที่ปลดหนี้หมด (นับจากเดือนนี้)').font=BLK;dp.cell(30,3,'=IF(ISNUMBER(C26),EDATE(TODAY(),C26),"-")').number_format='mmm yyyy';dp.cell(30,3).border=BOX;dp.cell(30,3).font=BLK
dp.cell(32,1,'ตารางจำลองรายเดือน (ด้านล่าง): เดือนที่ 1 = เดือนแรกที่เริ่มจ่ายตามแผน').font=f(size=10,italic=True,color='555555')
# sim columns: A month, B total balance end, C remaining budget unused, D interest this month, E paid this month,
# per debt k (0..5): bal_end col=6+k*4 ; interest col=7+k*4 ; pay col=8+k*4 ; extra col=9+k*4
hdr=['เดือนที่','ยอดหนี้รวมสิ้นเดือน','งบเหลือ (ไม่ใช้)','ดอกเบี้ยเดือนนี้','จ่ายรวมเดือนนี้']
for k in range(6):hdr+= [f'ก้อน {k+1} ยอดสิ้นเดือน',f'ก้อน {k+1} ดอกเบี้ย',f'ก้อน {k+1} จ่ายขั้นต่ำ',f'ก้อน {k+1} จ่ายเพิ่ม']
head(dp,S0,hdr)
for m in range(1,N+1):
    r=S0+m;dp.cell(r,1,m).font=BLK
    prevrow=r-1
    for k in range(6):
        cb=6+k*4;ci=7+k*4;cm=8+k*4;ce=9+k*4
        bprev=f'$C${18+k}' if m==1 else f'{CL(cb)}{prevrow}'
        dp.cell(r,ci,f'=ROUND({bprev}*$E${18+k}/12,2)')
        dp.cell(r,cm,f'=MIN($D${18+k},{bprev}+{CL(ci)}{r})')
        # remaining budget before this debt's extra: budget - all mins - extras of earlier debts
        allmins='+'.join(f'{CL(8+j*4)}{r}' for j in range(6))
        prevextra=('+'+'+'.join(f'{CL(9+j*4)}{r}' for j in range(k))) if k>0 else ''
        dp.cell(r,ce,f'=MAX(0,MIN($C$14-({allmins}){prevextra if False else ""}-{("("+"+".join(f"{CL(9+j*4)}{r}" for j in range(k)))+")" if k>0 else "0"},{bprev}+{CL(ci)}{r}-{CL(cm)}{r}))')
        dp.cell(r,cb,f'=ROUND({bprev}+{CL(ci)}{r}-{CL(cm)}{r}-{CL(ce)}{r},2)')
    dp.cell(r,2,'='+'+'.join(f'{CL(6+k*4)}{r}' for k in range(6)))
    dp.cell(r,4,'='+'+'.join(f'{CL(7+k*4)}{r}' for k in range(6)))
    dp.cell(r,5,'='+'+'.join(f'{CL(8+k*4)}{r}+{CL(9+k*4)}{r}' for k in range(6)))
    dp.cell(r,3,f'=$C$14-E{r}')
    for cc in range(1,30):
        c=dp.cell(r,cc);c.font=BLK;c.number_format=THB2 if cc>1 else '0'
for col,w in zip(range(1,30),[8,22,18,18,14,16]+[14]*23):dp.column_dimensions[CL(col)].width=w
dp.column_dimensions['M'].width=14
dp.freeze_panes=None

# ---------------- 4 เป้าหมายเก็บเงิน
gs=wb.create_sheet('เป้าหมายเก็บเงิน')
title(gs,'เป้าหมายเก็บเงิน','ใส่เป้าหมายและระยะเวลา ไฟล์จะคำนวณว่าต้องเก็บเดือนละ/สัปดาห์ละ/วันละเท่าไหร่ (ตัวอย่างในไฟล์: 300,000 บาทใน 36 เดือน ลองเปลี่ยนเป็น 1,000,000 ดู จะได้ ≈ 913 บาท/วัน)',6)
for col,w in zip('ABCDEF',[36,18,16,16,16,4]):gs.column_dimensions[col].width=w
lab=[('เป้าหมาย (บาท)',300000,THB),('ระยะเวลา (เดือน)',36,'0'),('มีเงินเก็บอยู่แล้ว (บาท)',0,THB),('ผลตอบแทนต่อปีที่คาดหวัง (ใส่ 0% ถ้าไม่รวม)',0.0,PCT)]
for i,(n,v,fm) in enumerate(lab):
    r=4+i;gs.cell(r,1,n).font=BLK;gs.cell(r,1).border=BOX;inp(gs.cell(r,2,v),fm)
gs.cell(8,1,'ผลตอบแทนเป็นสมมติฐานที่คุณกรอกเอง ไม่ได้รับประกัน ถ้าไม่แน่ใจให้ใช้ 0%').font=f(size=10,italic=True,color='C0392B')
gs.cell(10,1,'ต้องเก็บเดือนละ (บาท)').font=BOLD;gs.cell(10,1).border=BOX
calc(gs.cell(10,2,'=IF(B5<=0,0,IF(B7=0,MAX(0,B4-B6)/B5,MAX(0,-PMT(B7/12,B5,-B6,B4))))'),THB,True)
gs.cell(11,1,'ต้องเก็บสัปดาห์ละ (บาท)').font=BLK;gs.cell(11,1).border=BOX;calc(gs.cell(11,2,'=B10*12/52'),THB)
gs.cell(12,1,'ต้องเก็บวันละ (บาท)').font=BLK;gs.cell(12,1).border=BOX;calc(gs.cell(12,2,'=B10*12/365'),THB)
gs.cell(13,1,'งบเก็บเงินที่คุณแบ่งไว้ (จากแท็บ งบรายเดือน)').font=BLK;gs.cell(13,1).border=BOX;c=gs.cell(13,2,"='งบรายเดือน'!B27");c.font=GRN;c.border=BOX;c.number_format=THB
gs.cell(14,1,'ผล').font=BOLD;gs.cell(14,1).border=BOX
gs.cell(14,2,'=IF(B13>=B10,"พอ ✓ เก็บถึงเป้าตามเวลา","ยังขาดเดือนละ "&TEXT(B10-B13,"#,##0")&" บาท")').font=BOLD
gs.cell(15,1,'ถ้าเก็บตามงบที่แบ่งไว้ จะถึงเป้าในอีก (เดือน)').font=BLK;gs.cell(15,1).border=BOX
calc(gs.cell(15,2,'=IF(B13<=0,"-",IF(B7=0,IF(B4<=B6,0,ROUNDUP((B4-B6)/B13,0)),IFERROR(ROUNDUP(NPER(B7/12,-B13,-B6,B4),0),"-")))'),'0')
gs.cell(17,1,'วิธีคิด: ไม่รวมผลตอบแทน = (เป้าหมาย − เงินที่มี) ÷ จำนวนเดือน | รายวัน = รายเดือน × 12 ÷ 365 | ตัวเลขเป็นการประมาณ ไม่รวมภาษี/เงินเฟ้อ').font=f(size=10,italic=True,color='555555')
head(gs,19,['เดือนที่','ยอดสะสม (บาท)','% ถึงเป้า'])
for n in range(1,61):
    r=19+n;gs.cell(r,1,n).font=BLK;gs.cell(r,1).border=BOX
    gs.cell(r,2,f'=IF(A{r}>$B$5,"",IF($B$7=0,$B$6+$B$10*A{r},$B$6*(1+$B$7/12)^A{r}+$B$10*((1+$B$7/12)^A{r}-1)/($B$7/12)))').number_format=THB
    gs.cell(r,3,f'=IF(B{r}="","",IF($B$4=0,0,B{r}/$B$4))').number_format=PCT
    for cc in (2,3):gs.cell(r,cc).font=BLK;gs.cell(r,cc).border=BOX

# ---------------- 5 คำสั่ง AI
ai=wb.create_sheet('คำสั่ง AI')
title(ai,'คำสั่ง AI ช่วยตรวจและปรับแผน','คัดลอกข้อความในช่องไปวางใน ChatGPT / Claude แก้ข้อความใน [ ] ก่อนส่ง ห้ามใส่เลขบัญชีหรือรหัสผ่าน ให้ AI ตรวจตัวเลขแล้วตรวจซ้ำเองด้วย',3)
ai.column_dimensions['A'].width=30;ai.column_dimensions['B'].width=110
P=[('1. ตารางเก็บเงินตามเป้าหมาย','คุณคือผู้ช่วยวางแผนการเงินส่วนบุคคล ฉันอยากเก็บเงิน [จำนวนเงิน] บาท ภายใน [ระยะเวลา] รายได้ต่อเดือน [รายได้] บาท รายจ่ายประจำ [รายจ่าย] บาท ช่วยคำนวณว่าต้องเก็บเดือนละกี่บาท และทำตารางรายเดือน: เดือนที่ | ยอดที่เก็บ | ยอดสะสม | % ถึงเป้า แสดงวิธีคิดทุกขั้นตอน และระบุว่าไม่ได้รวมดอกเบี้ยและเงินเฟ้อ'),
('2. เช็กว่าเป้าหมายเป็นไปได้ไหม','จากรายได้ [รายได้] และรายจ่าย [รายจ่าย] ต่อเดือน เป้าหมายเก็บ [จำนวนเงิน] ใน [ระยะเวลา] เป็นไปได้จริงไหม ถ้าไม่ได้ ช่วยเสนอ 3 ทางเลือก: ขยายเวลา ลดเป้า หรือหารายได้เพิ่ม พร้อมตัวเลขของแต่ละทาง'),
('3. หาจุดรั่วในรายจ่าย','นี่คือรายจ่ายของฉันเดือนที่แล้ว: [วางรายการ] จัดหมวด จำเป็น/ควรมี/ฟุ่มเฟือย แล้วชี้ 3 รายการที่ลดได้มากสุด บอกตัวเลขที่ประหยัดได้ต่อเดือนและต่อปี อย่าตัดสินหรือตำหนิ'),
('4. จัดลำดับปลดหนี้','ฉันมีหนี้ดังนี้: [ชื่อหนี้ | ยอดคงเหลือ | ดอกเบี้ยต่อปี | ค่างวดขั้นต่ำ] และมีเงินเพิ่มปลดหนี้ได้เดือนละ [จำนวน] บาท เปรียบเทียบวิธี "ดอกเบี้ยสูงสุดก่อน" กับ "ยอดน้อยสุดก่อน" ว่าแต่ละวิธีหมดหนี้กี่เดือนและจ่ายดอกเบี้ยรวมเท่าไหร่ แสดงวิธีคิด'),
('5. ตรวจตัวเลขของ AI','ตรวจแผนที่คุณเพิ่งทำให้ฉัน: คำนวณใหม่ทีละขั้น บอกสมมติฐานที่ใช้ ชี้จุดที่อาจผิดหรือต้องตรวจเอง และถ้าไม่แน่ใจให้บอกว่าไม่แน่ใจ'),
('6. แผนเก็บเงินแบบวันละนิด','ฉันอยากเก็บเงิน [จำนวนเงิน] ใน [ระยะเวลา] ช่วยแปลงเป็นเป้ารายวันและรายสัปดาห์ เสนอกฎง่ายๆ 3 ข้อที่ทำตามได้จริง เช่น เก็บทันทีที่เงินเข้า'),
('7. เตรียมคุยกับเจ้าหนี้','ฉันผ่อนหนี้ [ประเภท] ไม่ไหว อยากขอปรับโครงสร้างหนี้หรือขอลดค่างวด ช่วยร่างข้อความสุภาพถึงเจ้าหนี้ และรายการเอกสารที่ควรเตรียม โดยไม่สัญญาในสิ่งที่ฉันจ่ายไม่ได้')]
head(ai,4,['หัวข้อ','คำสั่ง (คัดลอกทั้งช่อง)'])
for i,(t,x) in enumerate(P):
    r=5+i;a=ai.cell(r,1,t);a.font=BOLD;a.alignment=Alignment(wrap_text=True,vertical='top');a.border=BOX
    b=ai.cell(r,2,x);b.font=BLK;b.alignment=Alignment(wrap_text=True,vertical='top');b.border=BOX;ai.row_dimensions[r].height=62
wb.save('money-planner.xlsx');print('saved')
